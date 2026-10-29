import os
import docx

folder = r"practical_tests\Template - Progetto PM-20260227\Template - Progetto PM\Template_Progetto_PM"
files = sorted(
    [f for f in os.listdir(folder) if f.endswith(".docx")],
    key=lambda x: int(x.split()[0]) if x.split()[0].isdigit() else 999
)

print(f"Total docx templates: {len(files)}")
template_catalog = []

for f in files:
    path = os.path.join(folder, f)
    try:
        doc = docx.Document(path)
        tables = len(doc.tables)
        paras = len([p for p in doc.paragraphs if p.text.strip()])
        headers = []
        labels_found = []
        if tables > 0:
            for t in doc.tables[:3]:
                if len(t.rows) > 0:
                    r0 = [c.text.strip().replace("\n", " ") for c in t.rows[0].cells]
                    u = []
                    for h in r0:
                        if h and h not in u:
                            u.append(h)
                    headers.append("/".join(u[:3]))
                # Find some prominent labels in column 0
                for r in t.rows[:10]:
                    if r.cells:
                        txt = r.cells[0].text.strip().replace("\n", " ")
                        if txt and txt not in labels_found and len(txt) < 30:
                            labels_found.append(txt)
        info = {
            "file": f,
            "tables": tables,
            "paras": paras,
            "headers": headers,
            "labels": labels_found[:5]
        }
        template_catalog.append(info)
        h_str = " ; ".join(headers)[:35]
        l_str = ", ".join(labels_found[:3])
        print(f"{f:48} | T:{tables:2} | P:{paras:2} | H: {h_str:35} | Labels: {l_str}")
    except Exception as e:
        print(f"{f:48} | ERROR: {e}")
